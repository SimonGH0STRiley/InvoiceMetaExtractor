"""Export InvoiceRecord list to Excel (summary + exception list)."""



from __future__ import annotations



from pathlib import Path



from openpyxl import Workbook

from openpyxl.styles import Font, PatternFill



from core.models import InvoiceRecord



HEADERS = [

    "文件名",

    "发票类型",

    "发票号码",

    "开票日期",

    "销售方名称",

    "税目",

    "数量",

    "金额(不含税)",

    "税额",

    "税率",

    "校验",

]



EXCEPTION_HEADERS = HEADERS + ["源文件路径"]



_ISSUE_STATUSES = frozenset({"failed", "partial"})

_NUMBER_FORMAT = "#,##0.##"



def _validation_status_text(r: InvoiceRecord) -> str:

    if r.warnings or r.status != "success":

        return "校验失败"

    return "校验成功"





def _row_from_record(r: InvoiceRecord, *, include_path: bool = False) -> list:

    row = [

        Path(r.source_file).name,

        r.invoice_type,

        r.invoice_number,

        r.issue_date,

        r.seller_name,

        r.tax_items,

        r.quantity,

        r.amount,

        r.tax_amount,

        r.tax_rate,

        _validation_status_text(r),

    ]

    if include_path:

        row.append(r.source_file)

    return row





def _is_exception_record(r: InvoiceRecord) -> bool:

    if r.status in _ISSUE_STATUSES:

        return True

    return bool(r.warnings)





def _write_sheet_header(ws, headers: list[str]) -> None:

    ws.append(headers)

    for cell in ws[1]:

        cell.font = Font(bold=True)

        cell.fill = PatternFill("solid", fgColor="EEF2F7")





def _format_number_columns(ws) -> None:

    for col in (8, 9):  # 金额(不含税), 税额

        for row in range(2, ws.max_row + 1):

            ws.cell(row=row, column=col).number_format = _NUMBER_FORMAT





def export_records(records: list[InvoiceRecord], output_path: Path) -> Path:

    """Write summary sheet and exception list sheet."""

    wb = Workbook()

    ws_main = wb.active

    ws_main.title = "发票提取结果"

    _write_sheet_header(ws_main, HEADERS)



    exceptions: list[InvoiceRecord] = []

    for r in records:

        ws_main.append(_row_from_record(r))

        if _is_exception_record(r):

            exceptions.append(r)



    ws_exc = wb.create_sheet("异常清单")

    _write_sheet_header(ws_exc, EXCEPTION_HEADERS)

    for r in exceptions:

        ws_exc.append(_row_from_record(r, include_path=True))



    if not exceptions:

        ws_exc.append(["", "", "无异常记录"] + [""] * (len(EXCEPTION_HEADERS) - 3))


    _format_number_columns(ws_main)

    _format_number_columns(ws_exc)



    output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb.save(output_path)

    return output_path


def export_table(headers: list[str], rows: list[list], output_path: Path, *, sheet_name: str) -> Path:

    """Write a single table to an Excel workbook."""

    wb = Workbook()

    ws = wb.active

    ws.title = (sheet_name or "导出结果")[:31]

    _write_sheet_header(ws, headers)

    for row in rows:

        ws.append(row)

    for col in range(1, len(headers) + 1):

        max_length = max(

            len(str(ws.cell(row=row_index, column=col).value or ""))

            for row_index in range(1, ws.max_row + 1)

        )

        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(max(max_length + 2, 10), 36)

    output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb.save(output_path)

    return output_path


